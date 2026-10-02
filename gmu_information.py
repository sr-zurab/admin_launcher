import sys
import json
import requests
import urllib3
from PyQt6.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QLineEdit,
    QPushButton, QLabel, QTextEdit, QMessageBox,
    QDialog, QTableWidget, QTableWidgetItem, QHBoxLayout,
    QHeaderView
)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QIcon
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import sys
import os

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

API_BASE = "https://bus.gov.ru/public-rest/api/7710568760-Generalinformation"
BROWSER_HEADERS = {
    "User-Agent": "Mozilla/5.0",
    "Accept": "application/json"
}

FIXED_BIC = "019690001"
FIXED_RS = "03234643967010009400"
DATA_FILENAME = "data.json"


class BusGovApp(QWidget):
    def __init__(self):
        super().__init__()
        ICON_PATH = os.path.join(os.path.dirname(__file__), "list_org.ico")
        if os.path.exists(ICON_PATH):
            try:
                self.setWindowIcon(QIcon(ICON_PATH))
            except Exception:
                # если что-то пошло не так — пропускаем, приложение продолжит работать
                pass
        self.setWindowTitle("Информация об учреждении (bus.gov.ru)")
        self.resize(1000, 920)

        self.current_data = None

        main_layout = QVBoxLayout(self)

        # Ввод ИНН / КПП
        inn_layout = QHBoxLayout()
        self.inn_input = QLineEdit()
        self.inn_input.setPlaceholderText("ИНН (10 или 12 цифр)")
        self.kpp_input = QLineEdit()
        self.kpp_input.setPlaceholderText("КПП (9 цифр) — опционально")
        inn_layout.addWidget(QLabel("ИНН:"))
        inn_layout.addWidget(self.inn_input)
        inn_layout.addWidget(QLabel("КПП:"))
        inn_layout.addWidget(self.kpp_input)
        main_layout.addLayout(inn_layout)

        # Кнопка поиска
        self.search_button = QPushButton("Найти по ИНН/КПП")
        self.search_button.clicked.connect(self.search_org)
        main_layout.addWidget(self.search_button)

        # Результат / реквизиты
        main_layout.addWidget(QLabel("Реквизиты:"))
        self.result_area = QTextEdit()
        self.result_area.setReadOnly(True)
        main_layout.addWidget(self.result_area, stretch=1)

        # Ключ хранения
        key_layout = QHBoxLayout()
        self.key_input = QLineEdit()
        self.key_input.setPlaceholderText("Ключ хранения (название организации)")
        key_layout.addWidget(QLabel("Ключ:"))
        key_layout.addWidget(self.key_input)
        main_layout.addLayout(key_layout)

        # Кнопки операций (строка 1)
        btn_layout = QHBoxLayout()
        self.save_button = QPushButton("Сохранить")
        self.save_button.clicked.connect(self.save_data)
        self.load_button = QPushButton("Загрузить")
        self.load_button.clicked.connect(self.load_data)
        self.update_button = QPushButton("Обновить из API по ключу")
        self.update_button.clicked.connect(self.update_data)
        self.show_table_button = QPushButton("Показать таблицу организаций")
        self.show_table_button.clicked.connect(self.show_organizations_table)
        btn_layout.addWidget(self.save_button)
        btn_layout.addWidget(self.load_button)
        btn_layout.addWidget(self.update_button)
        btn_layout.addWidget(self.show_table_button)
        main_layout.addLayout(btn_layout)

        # Кнопки операций (строка 2)
        btn_layout2 = QHBoxLayout()
        self.delete_button = QPushButton("Удалить организацию")
        self.delete_button.clicked.connect(self.delete_organization)
        self.export_button = QPushButton("Экспорт в Excel")
        self.export_button.clicked.connect(self.export_all_to_excel)
        btn_layout2.addWidget(self.delete_button)
        btn_layout2.addWidget(self.export_button)
        main_layout.addLayout(btn_layout2)

    # --------------------------
    # Сетевой запрос к API
    # --------------------------
    def api_get(self, inn: str, kpp: str | None = None) -> requests.Response:
        url = f"{API_BASE}?inn={inn}" + (f"&kpp={kpp}" if kpp else "")
        return requests.get(url, headers=BROWSER_HEADERS, timeout=20, verify=False)

    # --------------------------
    # Поиск по ИНН/КПП
    # --------------------------
    def search_org(self):
        inn = self.inn_input.text().strip()
        kpp = self.kpp_input.text().strip() or None

        if not inn.isdigit() or len(inn) not in (10, 12):
            QMessageBox.warning(self, "Ошибка", "Введите корректный ИНН (10 или 12 цифр).")
            return
        if kpp and (not kpp.isdigit() or len(kpp) != 9):
            QMessageBox.warning(self, "Ошибка", "КПП должен быть 9 цифр или оставьте поле пустым.")
            return

        try:
            resp = self.api_get(inn, kpp)
        except Exception as e:
            QMessageBox.critical(self, "Ошибка сети", f"Ошибка запроса: {e}")
            return

        if resp.status_code == 200:
            try:
                data = resp.json()
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Не удалось распарсить JSON: {e}")
                return
            self.current_data = self.extract_fields(data)
            self.display_requisites()
        elif resp.status_code == 404:
            QMessageBox.information(self, "Не найдено", "Организация не найдена (404).")
        else:
            QMessageBox.warning(self, "Ошибка", f"HTTP {resp.status_code}")

    # --------------------------
    # Разбор ответа API
    # --------------------------
    def extract_fields(self, data: dict) -> dict:
        try:
            org = data["pagedContent"]["content"][0]["body"]["position"]
        except Exception:
            return {"initiator": {}, "main": {}, "other": {}}

        initiator = org.get("initiator", {}) or {}
        main = org.get("main", {}) or {}
        other = org.get("other", {}) or {}

        return {"initiator": initiator, "main": main, "other": other}

    # --------------------------
    # Построение Л/С (с сохранением всех символов остатка)
    # --------------------------
    def build_ls(self, reg_num: str) -> str:
        """
        Правило: Л/С = "20946" + regNum[3:] + "0"
        Тут важно: regNum рассматривается как строка, мы удаляем первые 3 символа
        (любой природы), берём остаток как есть и подставляем.
        Пример: "963Ц8835" -> остаток "Ц8835" -> "20946Ц88350"
        """
        if not reg_num:
            return ""
        s = str(reg_num)
        if len(s) <= 3:
            return ""
        remainder = s[3:]
        return "20946" + remainder + "0"

    # --------------------------
    # Красивый вывод классификаторов (каждое поле в новой строке)
    # --------------------------
    def pretty_classifier_lines(self, classifier: dict) -> list:
        """Возвращает список строк с отформатированными классификаторами на русском."""
        lines = []
        if not classifier or not isinstance(classifier, dict):
            return lines
        okfs = classifier.get("okfs")
        if okfs:
            code = okfs.get("code", "")
            name = okfs.get("name", "")
            lines.append(f"  ОКФС: {code}" + (f" — {name}" if name else ""))
        okopf = classifier.get("okopf")
        if okopf:
            code = okopf.get("code", "")
            name = okopf.get("name", "")
            lines.append(f"  ОКОПФ: {code}" + (f" — {name}" if name else ""))
        okpo = classifier.get("okpo")
        if okpo:
            lines.append(f"  ОКПО: {okpo}")
        oktmo = classifier.get("oktmo")
        if oktmo:
            code = oktmo.get("code", "")
            name = oktmo.get("name", "")
            lines.append(f"  ОКТМО: {code}" + (f" — {name}" if name else ""))
        okogu = classifier.get("okogu")
        if okogu:
            code = okogu.get("code", "")
            name = okogu.get("name", "")
            lines.append(f"  ОКОГУ: {code}" + (f" — {name}" if name else ""))
        okved = classifier.get("okved") or []
        if okved:
            for v in okved:
                code = v.get("code", "")
                name = v.get("name", "")
                typ = v.get("type", "")
                lines.append(f"  ОКВЭД: {code}" + (f" — {name}" if name else "") + (f" ({typ})" if typ else ""))
        return lines

    # --------------------------
    # Формирование текста реквизитов (каждое поле отдельной строкой "Наименование: значение")
    # --------------------------
    def format_requisites_text(self, data: dict) -> str:
        initiator = data.get("initiator", {}) or {}
        main = data.get("main", {}) or {}
        other = data.get("other", {}) or {}

        reg_num = initiator.get("regNum", "") or ""
        ls = self.build_ls(reg_num)

        # Адрес (комплексный) — собираем отдельными частями
        complex_address = main.get("complexAddress", {}) or {}
        address_obj = complex_address.get("address", {}) or {}
        addr_parts = []
        if address_obj.get("zip"):
            addr_parts.append(address_obj.get("zip"))
        subj = address_obj.get("subject", {}) or {}
        if subj.get("name"):
            addr_parts.append(subj.get("name"))
        city = address_obj.get("city", {}) or {}
        if city.get("name"):
            addr_parts.append(city.get("name"))
        if address_obj.get("street", {}) and address_obj["street"].get("name"):
            addr_parts.append(address_obj["street"].get("name"))
        if address_obj.get("businessPlace"):
            addr_parts.append(address_obj.get("businessPlace"))
        if address_obj.get("building"):
            addr_parts.append(address_obj.get("building"))
        address = ", ".join([p for p in addr_parts if p])

        # Классификатор — получаем список строк
        classifier = main.get("classifier", {}) or {}
        classifier_lines = self.pretty_classifier_lines(classifier)

        # Руководитель
        chiefs = other.get("chief", []) or []
        chief_line = ""
        if chiefs:
            c = chiefs[0]
            name_parts = [c.get("lastName",""), c.get("firstName",""), c.get("middleName","")]
            name = " ".join([p for p in name_parts if p]).strip()
            pos = c.get("position","")
            chief_line = name + (f" ({pos})" if pos else "")

        # Учредитель
        founders = other.get("founder", []) or []
        founder_line = ""
        if founders:
            f = founders[0]
            founder_line = f.get("fullName","") or ""

        # Собираем текст: каждая строка — "  Наименование: значение"
        lines = []
        lines.append("📌 Инициатор:")
        lines.append(f"  Реестровый номер: {reg_num}")
        lines.append(f"  Полное наименование: {initiator.get('fullName','')}")
        lines.append(f"  Краткое наименование: {main.get('shortName','')}")
        lines.append(f"  ИНН: {initiator.get('inn','')}")
        lines.append(f"  КПП: {initiator.get('kpp','')}")
        lines.append(f"  ОГРН: {main.get('ogrn','')}")
        lines.append(f"  Тип организации: {main.get('orgType','')}")

        # Вставляем классификаторы построчно (если есть)
        if classifier_lines:
            lines.append("  Классификаторы:")
            lines.extend(classifier_lines)

        # Адрес
        if address:
            lines.append(f"  Адрес: {address}")

        # Банковские реквизиты
        lines.append("")
        lines.append("💳 Банковские реквизиты:")
        lines.append(f"  Л/С: {ls}")
        lines.append(f"  БИК: {FIXED_BIC}")
        lines.append(f"  Р/С: {FIXED_RS}")

        # Руководитель
        if chief_line:
            lines.append("")
            lines.append("👤 Руководитель:")
            lines.append(f"  {chief_line}")

        # Учредитель
        if founder_line:
            lines.append("")
            lines.append("🏛 Учредитель:")
            lines.append(f"  {founder_line}")

        return "\n".join(lines)

    # --------------------------
    # Отобразить реквизиты
    # --------------------------
    def display_requisites(self):
        if not self.current_data:
            self.result_area.setPlainText("Данные отсутствуют.")
            return
        txt = self.format_requisites_text(self.current_data)
        self.result_area.setPlainText(txt)

    # --------------------------
    # Сохранение в data.json
    # --------------------------
    def save_data(self):
        if not self.current_data:
            QMessageBox.warning(self, "Ошибка", "Нет данных для сохранения.")
            return
        key = self.key_input.text().strip()
        if not key:
            QMessageBox.warning(self, "Ошибка", "Введите ключ хранения (название организации).")
            return
        try:
            try:
                with open(DATA_FILENAME, "r", encoding="utf-8") as f:
                    db = json.load(f)
            except FileNotFoundError:
                db = {}
            db[key] = self.current_data
            with open(DATA_FILENAME, "w", encoding="utf-8") as f:
                json.dump(db, f, ensure_ascii=False, indent=2)
            QMessageBox.information(self, "Сохранено", f"Организация сохранена под ключом '{key}'.")
        except Exception as e:
            QMessageBox.critical(self, "Ошибка сохранения", str(e))

    # --------------------------
    # Загрузка по ключу
    # --------------------------
    def load_data(self):
        key = self.key_input.text().strip()
        if not key:
            QMessageBox.warning(self, "Ошибка", "Введите ключ для загрузки.")
            return
        try:
            with open(DATA_FILENAME, "r", encoding="utf-8") as f:
                db = json.load(f)
        except FileNotFoundError:
            QMessageBox.warning(self, "Ошибка", f"Файл базы '{DATA_FILENAME}' не найден.")
            return
        entry = db.get(key)
        if not entry:
            QMessageBox.warning(self, "Ошибка", f"Ключ '{key}' не найден в базе.")
            return
        self.current_data = entry
        self.display_requisites()
        QMessageBox.information(self, "Загружено", f"Данные под ключом '{key}' загружены.")

    # --------------------------
    # Обновление по ключу (из API)
    # --------------------------
    def update_data(self):
        key = self.key_input.text().strip()
        if not key:
            QMessageBox.warning(self, "Ошибка", "Введите ключ для обновления.")
            return
        try:
            with open(DATA_FILENAME, "r", encoding="utf-8") as f:
                db = json.load(f)
        except FileNotFoundError:
            QMessageBox.warning(self, "Ошибка", f"Файл базы '{DATA_FILENAME}' не найден.")
            return
        entry = db.get(key)
        if not entry:
            QMessageBox.warning(self, "Ошибка", f"Ключ '{key}' не найден.")
            return
        inn = entry.get("initiator", {}).get("inn", "")
        kpp = entry.get("initiator", {}).get("kpp", "") or None
        if not inn:
            QMessageBox.warning(self, "Ошибка", "В записи нет ИНН для обновления.")
            return
        try:
            resp = self.api_get(inn, kpp)
        except Exception as e:
            QMessageBox.critical(self, "Ошибка сети", f"Ошибка запроса: {e}")
            return
        if resp.status_code == 200:
            try:
                new_data = self.extract_fields(resp.json())
            except Exception as e:
                QMessageBox.critical(self, "Ошибка", f"Не удалось распарсить JSON: {e}")
                return
            db[key] = new_data
            try:
                with open(DATA_FILENAME, "w", encoding="utf-8") as f:
                    json.dump(db, f, ensure_ascii=False, indent=2)
            except Exception as e:
                QMessageBox.critical(self, "Ошибка записи", str(e))
                return
            self.current_data = new_data
            self.display_requisites()
            QMessageBox.information(self, "Обновлено", f"Данные по ключу '{key}' обновлены из API.")
        elif resp.status_code == 404:
            QMessageBox.information(self, "Не найдено", "Организация с таким ИНН/КПП не найдена.")
        else:
            QMessageBox.warning(self, "Ошибка", f"HTTP {resp.status_code}")

    # --------------------------
    # Показать таблицу сохранённых организаций
    # --------------------------
    def show_organizations_table(self):
        try:
            with open(DATA_FILENAME, "r", encoding="utf-8") as f:
                db = json.load(f)
        except FileNotFoundError:
            QMessageBox.information(self, "Таблица", "База пуста или файл не найден.")
            return

        if not db:
            QMessageBox.information(self, "Таблица", "Нет сохранённых организаций.")
            return

        dialog = QDialog(self)
        dialog.setWindowTitle("Сохранённые организации")
        dialog.resize(1000, 600)
        layout = QVBoxLayout(dialog)

        table = QTableWidget()
        keys = list(db.keys())
        table.setRowCount(len(keys))
        table.setColumnCount(9)
        headers = ["Ключ", "Название", "ИНН", "КПП", "ОГРН", "Тип", "Адрес", "Руководитель", "Учредитель"]
        table.setHorizontalHeaderLabels(headers)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)

        for row, key in enumerate(keys):
            org = db[key]
            initiator = org.get("initiator", {}) or {}
            main = org.get("main", {}) or {}
            other = org.get("other", {}) or {}

            name = initiator.get("fullName") or main.get("shortName") or key
            inn = initiator.get("inn", "")
            kpp = initiator.get("kpp", "")
            ogrn = main.get("ogrn", "")
            org_type = main.get("orgType", "")

            # address
            addr_obj = main.get("complexAddress", {}).get("address", {}) or {}
            parts = []
            if addr_obj.get("zip"):
                parts.append(addr_obj.get("zip"))
            subj = addr_obj.get("subject", {}) or {}
            if subj.get("name"):
                parts.append(subj.get("name"))
            city = addr_obj.get("city", {}) or {}
            if city.get("name"):
                parts.append(city.get("name"))
            if addr_obj.get("street", {}) and addr_obj["street"].get("name"):
                parts.append(addr_obj["street"].get("name"))
            if addr_obj.get("building"):
                parts.append(addr_obj.get("building"))
            address = ", ".join([p for p in parts if p])

            chief_str = ""
            chiefs = other.get("chief", []) or []
            if chiefs:
                c = chiefs[0]
                name_parts = [c.get("lastName",""), c.get("firstName",""), c.get("middleName","")]
                chief_str = " ".join([p for p in name_parts if p]).strip()
                pos = c.get("position","")
                if pos:
                    chief_str += f" ({pos})"

            founder_str = ""
            founders = other.get("founder", []) or []
            if founders:
                founder_str = founders[0].get("fullName","") or ""

            table.setItem(row, 0, QTableWidgetItem(key))
            table.setItem(row, 1, QTableWidgetItem(name))
            table.setItem(row, 2, QTableWidgetItem(inn))
            table.setItem(row, 3, QTableWidgetItem(kpp))
            table.setItem(row, 4, QTableWidgetItem(ogrn))
            table.setItem(row, 5, QTableWidgetItem(org_type))
            table.setItem(row, 6, QTableWidgetItem(address))
            table.setItem(row, 7, QTableWidgetItem(chief_str))
            table.setItem(row, 8, QTableWidgetItem(founder_str))

        layout.addWidget(table)

        def on_double_click(item):
            sel_row = table.currentRow()
            if sel_row < 0:
                return
            selected_key = table.item(sel_row, 0).text()
            self.key_input.setText(selected_key)
            self.load_data()
            dialog.accept()

        table.itemDoubleClicked.connect(on_double_click)
        dialog.exec()

    # --------------------------
    # Удаление организации
    # --------------------------
    def delete_organization(self):
        key = self.key_input.text().strip()
        if not key:
            QMessageBox.warning(self, "Ошибка", "Введите ключ для удаления.")
            return
        try:
            with open(DATA_FILENAME, "r", encoding="utf-8") as f:
                db = json.load(f)
        except FileNotFoundError:
            QMessageBox.warning(self, "Ошибка", "Файл базы не найден.")
            return
        if key not in db:
            QMessageBox.warning(self, "Ошибка", f"Ключ '{key}' не найден.")
            return
        confirm = QMessageBox.question(self, "Подтвердите удаление", f"Удалить организацию '{key}'?", QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
        if confirm != QMessageBox.StandardButton.Yes:
            return
        try:
            del db[key]
            with open(DATA_FILENAME, "w", encoding="utf-8") as f:
                json.dump(db, f, ensure_ascii=False, indent=2)
            QMessageBox.information(self, "Удалено", f"Организация '{key}' удалена.")
            if self.current_data and self.key_input.text().strip() == key:
                self.current_data = None
                self.result_area.clear()
        except Exception as e:
            QMessageBox.critical(self, "Ошибка удаления", str(e))

    # --------------------------
    # Экспорт в Excel
    # --------------------------
    def export_all_to_excel(self):
        try:
            with open(DATA_FILENAME, "r", encoding="utf-8") as f:
                db = json.load(f)
        except FileNotFoundError:
            QMessageBox.warning(self, "Ошибка", "Файл базы не найден.")
            return
        if not db:
            QMessageBox.information(self, "Экспорт", "База пуста.")
            return

        filename = "organizations_export.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Организации"

        headers = ["Ключ", "Название", "ИНН", "КПП", "ОГРН", "Тип", "Адрес", "Руководитель", "Учредитель", "Л/С", "БИК", "Р/С", "ОКФС", "ОКОПФ", "ОКПО", "ОКТМО", "ОКВЭД"]
        ws.append(headers)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4F81BD")
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for key, org in db.items():
            initiator = org.get("initiator", {}) or {}
            main = org.get("main", {}) or {}
            other = org.get("other", {}) or {}

            name = initiator.get("fullName") or main.get("shortName") or key
            inn = initiator.get("inn", "")
            kpp = initiator.get("kpp", "")
            ogrn = main.get("ogrn", "")
            org_type = main.get("orgType", "")

            # address
            addr_obj = main.get("complexAddress", {}).get("address", {}) or {}
            parts = []
            if addr_obj.get("zip"):
                parts.append(addr_obj.get("zip"))
            subj = addr_obj.get("subject", {}) or {}
            if subj.get("name"):
                parts.append(subj.get("name"))
            city = addr_obj.get("city", {}) or {}
            if city.get("name"):
                parts.append(city.get("name"))
            if addr_obj.get("street", {}) and addr_obj["street"].get("name"):
                parts.append(addr_obj["street"].get("name"))
            if addr_obj.get("building"):
                parts.append(addr_obj.get("building"))
            address = ", ".join([p for p in parts if p])

            chief_str = ""
            chiefs = other.get("chief", []) or []
            if chiefs:
                c = chiefs[0]
                name_parts = [c.get("lastName", ""), c.get("firstName", ""), c.get("middleName", "")]
                pos = c.get("position", "")
                chief_str = " ".join([p for p in name_parts if p]).strip() + (f" ({pos})" if pos else "")

            founder_str = ""
            founders = other.get("founder", []) or []
            if founders:
                founder_str = founders[0].get("fullName", "") or ""

            reg_num = initiator.get("regNum", "") or ""
            ls = self.build_ls(reg_num)

            # classifier breakdown
            classifier = main.get("classifier", {}) or {}
            okfs = classifier.get("okfs", {}) or {}
            okopf = classifier.get("okopf", {}) or {}
            okpo = classifier.get("okpo", "") or ""
            oktmo = classifier.get("oktmo", {}) or {}
            okved = classifier.get("okved", []) or []
            okved_txt = "; ".join([f"{v.get('code','')} — {v.get('name','')}" for v in okved])

            ws.append([
                key, name, inn, kpp, ogrn, org_type, address,
                chief_str, founder_str, ls, FIXED_BIC, FIXED_RS,
                okfs.get("code","") + (f" — {okfs.get('name','')}" if okfs.get("name") else ""),
                okopf.get("code","") + (f" — {okopf.get('name','')}" if okopf.get("name") else ""),
                okpo,
                oktmo.get("code","") + (f" — {oktmo.get('name','')}" if oktmo.get("name") else ""),
                okved_txt
            ])

        for column_cells in ws.columns:
            length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = min(60, length + 2)

        try:
            wb.save(filename)
            QMessageBox.information(self, "Экспорт", f"Экспорт завершён: {filename}")
        except Exception as e:
            QMessageBox.critical(self, "Ошибка экспорта", str(e))
